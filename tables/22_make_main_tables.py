"""Build three-line (黑白三线表) versions of Tables and Supplementary Tables.

Three-line table convention:
  - Top rule (1.5 pt black) at top of table
  - Middle rule (0.75 pt black) below header row
  - Bottom rule (1.5 pt black) at bottom of table
  - NO vertical lines, NO fill colours, NO interior horizontal lines
"""
import json, pandas as pd, numpy as np
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

def remove_all_borders(table):
    """Strip every border from every cell."""
    tbl = table._tbl
    tblPr = tbl.find(qn('w:tblPr'))
    if tblPr is None:
        tblPr = OxmlElement('w:tblPr')
        tbl.insert(0, tblPr)
    # Remove existing tblBorders
    existing = tblPr.find(qn('w:tblBorders'))
    if existing is not None:
        tblPr.remove(existing)
    tblBorders = OxmlElement('w:tblBorders')
    for edge in ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']:
        b = OxmlElement(f'w:{edge}')
        b.set(qn('w:val'), 'nil')
        tblBorders.append(b)
    tblPr.append(tblBorders)
    # Strip any direct cell borders
    for row in table.rows:
        for cell in row.cells:
            tcPr = cell._tc.get_or_add_tcPr()
            existing_cell = tcPr.find(qn('w:tcBorders'))
            if existing_cell is not None:
                tcPr.remove(existing_cell)
            tcBorders = OxmlElement('w:tcBorders')
            for edge in ['top', 'left', 'bottom', 'right']:
                b = OxmlElement(f'w:{edge}')
                b.set(qn('w:val'), 'nil')
                tcBorders.append(b)
            tcPr.append(tcBorders)

def set_row_border(table, row_idx, edge='bottom', sz=12):
    """Add a horizontal border to a specific row.
    edge='top' or 'bottom'; sz in 1/8 points (sz=12 → 1.5 pt; sz=6 → 0.75 pt)."""
    for cell in table.rows[row_idx].cells:
        tcPr = cell._tc.get_or_add_tcPr()
        tcBorders = tcPr.find(qn('w:tcBorders'))
        if tcBorders is None:
            tcBorders = OxmlElement('w:tcBorders')
            tcPr.append(tcBorders)
        # Remove existing edge if present
        existing = tcBorders.find(qn(f'w:{edge}'))
        if existing is not None:
            tcBorders.remove(existing)
        b = OxmlElement(f'w:{edge}')
        b.set(qn('w:val'), 'single')
        b.set(qn('w:sz'), str(sz))
        b.set(qn('w:color'), '000000')
        tcBorders.append(b)

def remove_cell_shading(cell):
    """Strip background fill."""
    tcPr = cell._tc.get_or_add_tcPr()
    existing = tcPr.find(qn('w:shd'))
    if existing is not None:
        tcPr.remove(existing)
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), 'auto')
    tcPr.append(shd)

def make_three_line_table(doc, rows_data, font_size=9):
    """Create a clean three-line table with no colour, no vertical/interior lines."""
    n_rows = len(rows_data); n_cols = len(rows_data[0])
    t = doc.add_table(rows=n_rows, cols=n_cols)
    t.style = 'Table Grid'   # has a definite cell box that we'll then strip
    t.alignment = WD_TABLE_ALIGNMENT.CENTER

    # Set content first
    for ri, row_data in enumerate(rows_data):
        for ci, txt in enumerate(row_data):
            cell = t.cell(ri, ci)
            cell.text = ""
            p = cell.paragraphs[0]
            run = p.add_run(str(txt))
            run.font.size = Pt(font_size)
            run.font.name = 'Times New Roman'
            if ri == 0:
                run.font.bold = True
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            remove_cell_shading(cell)

    # Strip all borders
    remove_all_borders(t)

    # Apply three rules:
    set_row_border(t, 0, edge='top',    sz=12)  # top rule 1.5 pt
    set_row_border(t, 0, edge='bottom', sz=6)   # mid rule 0.75 pt
    set_row_border(t, n_rows-1, edge='bottom', sz=12)  # bottom rule 1.5 pt

    return t

def add_caption(doc, text, size=10, space_before=Pt(14)):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = space_before
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(text)
    run.font.size = Pt(size); run.font.bold = True
    run.font.name = 'Times New Roman'
    return p

def add_title(doc, text, size=12):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(text)
    r.font.size = Pt(size); r.font.bold = True
    r.font.name = 'Times New Roman'
    return p

# =============================================================================
# MAIN TABLES — build new Tables_revised.docx as three-line tables
# =============================================================================
doc = Document()
section = doc.sections[0]
section.top_margin = Cm(2.0); section.bottom_margin = Cm(2.0)
section.left_margin = Cm(2.0); section.right_margin = Cm(2.0)
style = doc.styles['Normal']
style.font.name = 'Times New Roman'
style.font.size = Pt(10)

add_title(doc, "Tables — Revised manuscript ISCIENCE-D-26-03472", size=13)

# Table 1
add_caption(doc, "Table 1.  Cohort characteristics for the discovery, "
                  "external, independent-external, and tissue-boundary "
                  "validation cohorts.")
t1 = [
    ["Characteristic", "GSE18123 (GPL570)", "GSE18123 (GPL6244)", "GSE42133", "GSE25507"],
    ["Role", "Discovery / training", "External validation (whole blood)",
     "Independent external (leukocyte)", "Tissue-boundary exploration"],
    ["Platform", "Affymetrix HG-U133 Plus 2.0", "Affymetrix HuGene-1_0-ST",
     "Illumina HumanHT-12 v4", "Affymetrix HG-U133 Plus 2.0"],
    ["Tissue", "Peripheral whole blood", "Peripheral whole blood",
     "Peripheral leukocyte", "Lymphocyte-enriched mononuclear cells"],
    ["Total samples (n)", "99", "186", "147", "146"],
    ["ASD cases", "66", "104", "91", "82"],
    ["Typically developing", "33", "82", "56", "64"],
    ["Sex composition", "Male only (analytic restriction)",
     "Male 128, Female 58", "Male only (per accession metadata)",
     "Mixed; sex unrecoverable from data"],
    ["Age range", "Pediatric (≈2–14 y)", "Pediatric (≈2–14 y)",
     "Pediatric (≈1–4 y; Pierce cohort)", "Pediatric (≈3–15 y)"],
    ["Normalisation", "log₂(series matrix)", "log₂(series matrix)",
     "log₂(series matrix)", "log₂(series matrix)"],
    ["7 panel genes mappable", "7 / 7", "7 / 7",
     "6 / 7 (ERVK3-2 absent on platform)", "7 / 7"],
]
make_three_line_table(doc, t1)

# Table 2
add_caption(doc, "Table 2.  Locked 7-gene panel: coefficients, training-set "
                  "mean and SD on the log₂ scale, and selection frequency "
                  "across 500 outer folds of stability selection.")
t2 = [
    ["Gene", "Biotype", "Coefficient (logit)", "Train mean (log₂)",
     "Train SD (log₂)", "Direction", "Sel. freq."],
    ["KYNU", "Protein-coding", "+1.4064", "2.060", "1.266", "Up in ASD", "98.6 %"],
    ["CMYA5", "Protein-coding", "+0.9988", "1.815", "0.875", "Up in ASD", "78.0 %"],
    ["MAPK8IP1", "Protein-coding", "+0.9149", "2.506", "2.386", "Up in ASD", "71.6 %"],
    ["CES1", "Protein-coding", "+0.8085", "6.541", "0.863", "Up in ASD", "63.0 %"],
    ["KIAA0087", "lncRNA (poorly characterised)", "−0.7340", "2.264", "0.826",
     "Down in ASD", "82.4 %"],
    ["POU2AF1", "Protein-coding (B-cell associated)", "−0.5666", "3.238", "0.897",
     "Down in ASD", "62.0 %"],
    ["ERVK3-2", "Endogenous retroviral", "−0.4620", "2.368", "1.253",
     "Down in ASD", "62.0 %"],
    ["Intercept (β₀)", "—", "+1.0188", "—", "—", "—", "—"],
]
make_three_line_table(doc, t2)

# Table 3
add_caption(doc, "Table 3.  Discrimination and calibration of the 7-gene panel "
                  "across cohorts.")
t3 = [
    ["Cohort", "Validation tier", "Score / mode", "n (ASD / Ctrl)",
     "AUC", "95 % CI", "Brier"],
    ["GSE18123-GPL570", "Internal — apparent", "LP₇ (training fit)",
     "66 / 33", "0.942", "—", "0.097"],
    ["GSE18123-GPL570", "Internal — repeated nested CV",
     "5-fold × 100 repeats", "66 / 33", "0.762", "0.677–0.834", "—"],
    ["GSE18123-GPL570", "Internal — bootstrap (.632)",
     "B = 1 000, optimism-corrected", "66 / 33", "0.891", "—", "—"],
    ["GSE18123-GPL570", "Internal — permutation null",
     "5-fold CV, 1 000 perms", "66 / 33", "p = 0.001", "—", "—"],
    ["GSE18123-GPL6244", "External — platform-adapted",
     "LP₇ (panel-adapted z-score)", "104 / 82", "0.689", "0.608–0.761", "0.260"],
    ["GSE18123-GPL6244", "External — strict train scale",
     "LP₇ (training z-score)", "104 / 82", "0.695", "0.616–0.766", "0.441"],
    ["GSE18123-GPL6244", "External — males only",
     "LP₇ (platform-adapted)", "80 / 48", "0.704", "0.612–0.797", "0.225"],
    ["GSE18123-GPL6244", "External — females only",
     "LP₇ (platform-adapted)", "24 / 34", "0.640", "0.485–0.794", "0.343"],
    ["GSE18123-GPL6244", "External — after one-step recalibration",
     "α = −0.058, β = +0.304", "104 / 82", "0.689", "—", "0.221"],
    ["GSE18123-GPL6244", "External — permutation null",
     "10 000 label permutations", "104 / 82", "p < 0.0001", "—", "—"],
    ["GSE42133", "Independent external — 6-gene refit",
     "LP₆ (panel-adapted z-score)", "91 / 56", "0.420", "0.326–0.519", "0.379"],
    ["GSE42133", "Independent external — 7-gene zero-fill",
     "Missing gene zero-filled", "91 / 56", "0.421", "0.327–0.523", "0.373"],
    ["GSE42133", "Independent external — permutation null",
     "10 000 perms, two-sided", "91 / 56", "p = 0.12", "—", "—"],
    ["GSE25507", "Tissue boundary — adapted",
     "LP₇ on lymphocyte-enriched", "82 / 64", "0.464", "0.368–0.561", "—"],
    ["GSE25507", "Tissue boundary — permutation null",
     "10 000 perms, two-sided", "82 / 64", "p = 0.46", "—", "—"],
    ["GSE25507", "Tissue boundary — within-tissue refit",
     "Same 7 genes, re-trained", "82 / 64", "0.620", "—", "—"],
]
make_three_line_table(doc, t3)

# Table 4
add_caption(doc, "Table 4.  Pearson correlation of LP₇ (or LP₆ on GSE42133) "
                  "with literature-defined blood-cell-type marker scores in "
                  "each of the four cohorts.")
t4 = [
    ["Cell-type marker (Pearson r vs LP)",
     "GPL570 (whole blood)", "GPL6244 (whole blood)",
     "GSE42133 (leukocyte)", "GSE25507 (lymphocyte)",
     "Cross-cohort pattern"],
    ["Monocytes (8 markers)", "+0.567 ***", "+0.411 ***", "+0.508 ***", "+0.466 ***",
     "Consistent — the signature dimension"],
    ["Neutrophils (8 markers)", "+0.266 **", "+0.035", "+0.254 **", "+0.193 *",
     "Variable; weak"],
    ["Dendritic cells (4 markers)", "+0.397 ***", "+0.082", "+0.151", "−0.110",
     "Whole-blood-specific"],
    ["NK cells (6 markers)", "+0.329 ***", "+0.139", "−0.096", "+0.132",
     "Whole-blood-specific (weak)"],
    ["B cells (8 markers)", "−0.199 *", "−0.315 ***", "−0.368 ***", "−0.348 ***",
     "Consistently negative; stronger in non-whole-blood"],
    ["CD8 T cells (6 markers)", "+0.252 *", "−0.007", "−0.255 **", "−0.038",
     "Variable across platforms"],
    ["CD4 T cells (7 markers)", "+0.056", "−0.191 **", "−0.333 ***", "−0.232 **",
     "Negative in mixed-sex / non-whole-blood"],
    ["Erythrocytes (5 markers)", "+0.050", "+0.050", "+0.121", "+0.042",
     "Negligible"],
]
make_three_line_table(doc, t4)

# Footnote
foot = doc.add_paragraph()
foot.paragraph_format.space_before = Pt(6)
fr = foot.add_run("Significance: * p < 0.05, ** p < 0.01, *** p < 0.001 (uncorrected).")
fr.font.size = Pt(9); fr.font.italic = True; fr.font.name = 'Times New Roman'

doc.save('../Tables_revised.docx')
print("Main tables written")
